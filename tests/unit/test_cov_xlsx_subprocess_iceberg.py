from __future__ import annotations

import pickle
import subprocess
import sys
from pathlib import Path
from datetime import UTC, datetime
from unittest.mock import patch, MagicMock

import pytest
from openpyxl import Workbook

from sre.compiler import CompiledRulePackage
from sre.ioxls import DecisionTableImporter
from sre.ioxls.exporter import DecisionTableExporter
from sre.model import DecisionTable, HitPolicy
from sre.model.rule import active_set_hash, Rule, RuleDefinition, RuleFormat, new_rule_id
from sre.runtime.iceberg_store import IcebergLikeTable, UnknownSnapshotError
from sre.store import InMemoryRuleMetadataStore, UnknownRuleError


def _write_min_xlsx(
    p: Path,
    *,
    b2: str = "FIRST",
    cell_row8_col1: int | str = 1,
) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "RuleTable"
    ws["C1"] = "mytab"
    ws["B2"] = b2
    ws.cell(3, 1, "CONDITION")
    ws.cell(4, 1, "a")
    ws.cell(5, 1, "a")
    ws.cell(6, 1, "INT")
    ws.cell(7, 1, "==")
    ws.cell(3, 2, "ACTION")
    ws.cell(4, 2, "o")
    ws.cell(5, 2, "o")
    ws.cell(6, 2, "STRING")
    ws["A8"] = cell_row8_col1
    ws["B8"] = "x"
    ws["C8"] = 0
    wb.save(p)


def test_importer_happy() -> None:
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "t.xlsx"
        _write_min_xlsx(p, b2="COLLECT", cell_row8_col1=1)
        r = DecisionTableImporter.import_file(p)
        assert r.table and r.table.name == "mytab"
        assert not r.errors


def test_importer_bad_policy() -> None:
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "t2.xlsx"
        _write_min_xlsx(p, b2="NOTAPOLICY")
        r = DecisionTableImporter.import_file(p)
        assert r.table is None and r.errors


def test_importer_type_error_int_cell() -> None:
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "t3.xlsx"
        _write_min_xlsx(p, cell_row8_col1="notint")
        r = DecisionTableImporter.import_file(p)
        r.table is None
        assert r.errors


def test_importer_no_columns() -> None:
    import tempfile

    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "emptyc.xlsx"
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws["A1"] = "X"
        ws["B2"] = "FIRST"
        ws["C1"] = "n"
        wb.save(p)
        r = DecisionTableImporter.import_file(p)
        assert r.table is None and r.errors


def test_importer_float_input_bad_cell() -> None:
    import tempfile

    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "fl.xlsx"
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws["A1"] = "T"
        ws["C1"] = "f"
        ws["B2"] = "FIRST"
        ws.cell(3, 1, "CONDITION")
        ws.cell(4, 1, "a")
        ws.cell(5, 1, "a")
        ws.cell(6, 1, "FLOAT")
        ws.cell(3, 2, "ACTION")
        ws.cell(4, 2, "o")
        ws.cell(5, 2, "o")
        ws.cell(6, 2, "STRING")
        ws["A8"] = "nope"
        ws["B8"] = "x"
        wb.save(p)
        r = DecisionTableImporter.import_file(p)
        assert r.table is None
        assert r.errors


def test_importer_bad_priority_defaults_zero() -> None:
    import tempfile

    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "pri.xlsx"
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws["A1"] = "T"
        ws["C1"] = "t"
        ws["B2"] = "FIRST"
        ws.cell(3, 1, "CONDITION")
        ws.cell(4, 1, "a")
        ws.cell(5, 1, "a")
        ws.cell(6, 1, "INT")
        ws.cell(3, 2, "ACTION")
        ws.cell(4, 2, "o")
        ws.cell(5, 2, "o")
        ws.cell(6, 2, "STRING")
        ws["A8"] = 1
        ws["B8"] = "ok"
        ws.cell(8, 3, "not_int")
        wb.save(p)
        r = DecisionTableImporter.import_file(p)
        assert r.table is not None
        assert r.table.rows[0].priority == 0


def test_importer_output_cell_type_error_and_bad_priority() -> None:
    import tempfile

    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "out_err.xlsx"
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws["A1"] = "T"
        ws["C1"] = "t"
        ws["B2"] = "FIRST"
        ws.cell(3, 1, "CONDITION")
        ws.cell(4, 1, "a")
        ws.cell(5, 1, "a")
        ws.cell(6, 1, "INT")
        ws.cell(3, 2, "ACTION")
        ws.cell(4, 2, "o")
        ws.cell(5, 2, "o")
        ws.cell(6, 2, "INT")
        ws["A8"] = 1
        ws["B8"] = "badint"
        ws.cell(8, 3, "not_an_int")
        wb.save(p)
        r = DecisionTableImporter.import_file(p)
        assert r.table is None
        assert r.errors


def test_importer_bool_bad_cell() -> None:
    import tempfile

    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "bl.xlsx"
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws["A1"] = "T"
        ws["C1"] = "f"
        ws["B2"] = "FIRST"
        ws.cell(3, 1, "CONDITION")
        ws.cell(4, 1, "a")
        ws.cell(5, 1, "a")
        ws.cell(6, 1, "BOOL")
        ws.cell(3, 2, "ACTION")
        ws.cell(4, 2, "o")
        ws.cell(5, 2, "o")
        ws.cell(6, 2, "STRING")
        ws["A8"] = "maybe"
        ws["B8"] = "x"
        wb.save(p)
        r = DecisionTableImporter.import_file(p)
        assert r.table is None
        assert r.errors


@patch("sre.ioxls.importer.load_workbook")
def test_importer_no_active_sheet(mock_lw) -> None:
    wb = MagicMock()
    wb.active = None
    mock_lw.return_value = wb
    r = DecisionTableImporter.import_file("dummy.xlsx")
    assert r.table is None
    assert r.errors


@patch("sre.ioxls.exporter.Workbook")
def test_exporter_workbook_no_active(mock_wb) -> None:
    w = MagicMock()
    w.active = None
    mock_wb.return_value = w
    dt = DecisionTable("x", HitPolicy.FIRST, (), (), ())
    with pytest.raises(RuntimeError, match="no active"):
        DecisionTableExporter.export(dt, "out.xlsx")


def test_subprocess_module_main() -> None:
    r = subprocess.run(
        [sys.executable, "-m", "sre.tools.smoke_drl"],
        cwd=Path(__file__).resolve().parents[2],
        capture_output=True,
        text=True,
    )
    assert r.returncode == 0
    assert "rule" in r.stdout.lower() or "r1" in r.stdout


def test_smoke_drl_runpy_as_main(monkeypatch) -> None:
    import runpy
    from io import StringIO

    monkeypatch.setattr(sys, "argv", ["smoke_drl"], raising=False)
    buf = StringIO()
    old = sys.stdout
    try:
        sys.stdout = buf
        with pytest.raises(SystemExit) as ex:
            runpy.run_module("sre.tools.smoke_drl", run_name="__main__")
        assert ex.value.code == 0
    finally:
        sys.stdout = old
    s = buf.getvalue()
    assert "rule" in s.lower() or "r1" in s


def test_active_set_hash_fn() -> None:
    h = active_set_hash([("a", 1), ("a", 2)])
    assert len(h) == 64
    s = InMemoryRuleMetadataStore()
    t0 = datetime(2019, 1, 1, tzinfo=UTC)
    s.insert(
        Rule(
            new_rule_id(), "a", 0, "g", 0, t0, None, True,
            RuleDefinition("x", RuleFormat.DRL), None,
        )
    )
    s.active_set_version(t0)
    with pytest.raises(UnknownRuleError):
        s.get("z", 1)


def test_compiled_rule_package_reject_bad() -> None:
    with pytest.raises((TypeError, Exception)):
        CompiledRulePackage.deserialize(b"not a pkg")


def test_iceberg_delete_and_pickle() -> None:
    t = IcebergLikeTable("F", {"x": int, "f": int})
    t.append([{"x": 1, "f": 0}])
    t.append([{"x": 1, "f": 0}, {"x": 2, "f": 0}])
    t.delete_rows(lambda r: r.get("f") == 0)
    with pytest.raises(UnknownSnapshotError):
        t.snapshot(9_999_999)
    b = pickle.dumps(t)
    t2 = pickle.loads(b)
    assert t2.name == t.name
    assert t2.current_snapshot_id() == t2.current_snapshot_id()
    t_append = IcebergLikeTable("A", {"x": int}, append_only=True)
    t_append.append([{"x": 1}])
    with pytest.raises(ValueError):
        t_append.delete_rows(lambda _r: True)


def test_iceberg_setstate_backward_compat() -> None:
    t = IcebergLikeTable("B", {"x": int})
    t.__setstate__(("B", {"x": int}, {0: [{"x": 1}]}, 0))
    assert t.append_only is False
    assert t.snapshot(0)[0]["x"] == 1
