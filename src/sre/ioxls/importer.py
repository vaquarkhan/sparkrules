from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from sre.model.decision_table import (
    ColumnType,
    DecisionTable,
    HitPolicy,
    InputColumn,
    OutputColumn,
    Row,
)


class XlsxImportError(Exception):
    pass


class CellTypeError(XlsxImportError):
    def __init__(self, message: str, *, row: int, col: int) -> None:
        super().__init__(message)
        self.row = row
        self.col = col


@dataclass
class ImportResult:
    table: DecisionTable | None
    errors: list[CellTypeError] = field(default_factory=list)


def _parse_type(s: str) -> ColumnType:
    return ColumnType[str(s).strip().upper()]


def _coerce(val: Any, ct: ColumnType, row: int, col: int) -> Any:
    if val is None or val == "":
        return val
    if ct == ColumnType.STRING:
        return str(val)
    if ct == ColumnType.INT:
        try:
            return int(val)
        except (TypeError, ValueError) as e:
            raise CellTypeError(str(e), row=row, col=col) from e
    if ct == ColumnType.FLOAT:
        try:
            return float(val)
        except (TypeError, ValueError) as e:
            raise CellTypeError(str(e), row=row, col=col) from e
    if ct == ColumnType.BOOL:
        if isinstance(val, bool):
            return val
        s = str(val).lower()
        if s in ("true", "1", "yes"):
            return True
        if s in ("false", "0", "no"):
            return False
        raise CellTypeError(f"not bool: {val!r}", row=row, col=col)
    return val


class DecisionTableImporter:
    @staticmethod
    def import_file(path: str | Path) -> ImportResult:
        errors: list[CellTypeError] = []
        p = Path(path)
        wb = load_workbook(p, data_only=True)
        ws = wb.active
        if ws is None:
            return ImportResult(
                None, [CellTypeError("no active sheet", row=0, col=0)]
            )
        name = str(ws["C1"].value or "table")
        b2 = str(ws["B2"].value or "UNIQUE").strip().upper()
        try:
            hp = HitPolicy[b2]
        except KeyError:
            return ImportResult(
                None,
                [CellTypeError(f"invalid HIT_POLICY {b2!r}", row=2, col=2)],
            )
        c = 1
        in_cols: list[InputColumn] = []
        out_cols: list[OutputColumn] = []
        while True:
            kind = ws.cell(3, c).value
            if kind in (None, ""):
                break
            kind_s = str(kind).strip()
            if kind_s == "CONDITION":
                in_cols.append(
                    InputColumn(
                        name=str(ws.cell(4, c).value or ""),
                        field_ref=str(ws.cell(5, c).value or ""),
                        col_type=_parse_type(str(ws.cell(6, c).value or "STRING")),
                        operator=(
                            str(ws.cell(7, c).value)
                            if ws.cell(7, c).value not in (None, "")
                            else "=="
                        ),
                    )
                )
            elif kind_s == "ACTION":
                out_cols.append(
                    OutputColumn(
                        name=str(ws.cell(4, c).value or ""),
                        field_ref=str(ws.cell(5, c).value or ""),
                        col_type=_parse_type(str(ws.cell(6, c).value or "STRING")),
                    )
                )
            c += 1
        n_in, n_out = len(in_cols), len(out_cols)
        if n_in + n_out == 0:
            return ImportResult(
                None, [CellTypeError("no columns in sheet", row=3, col=0)]
            )
        rows: list[Row] = []
        r = 8
        while r < 100_000:
            parts = [
                ws.cell(r, j).value
                for j in range(1, n_in + n_out + 1)
            ]
            if all(x in (None, "") for x in parts):
                break
            cells: list[Any] = []
            for j in range(1, n_in + 1):
                try:
                    cells.append(
                        _coerce(
                            ws.cell(r, j).value, in_cols[j - 1].col_type, r, j
                        )
                    )
                except CellTypeError as e:
                    errors.append(e)
            for j in range(n_in + 1, n_in + n_out + 1):
                oi = j - n_in - 1
                try:
                    cells.append(
                        _coerce(
                            ws.cell(r, j).value, out_cols[oi].col_type, r, j
                        )
                    )
                except CellTypeError as e:
                    errors.append(e)
            pri_v = ws.cell(r, n_in + n_out + 1).value
            pr = 0
            if pri_v not in (None, ""):
                try:
                    pr = int(pri_v)
                except (TypeError, ValueError):
                    pr = 0
            rows.append(Row(tuple(cells), priority=pr))
            r += 1
        if errors:
            return ImportResult(None, errors)
        return ImportResult(
            DecisionTable(
                name=name,
                hit_policy=hp,
                input_columns=tuple(in_cols),
                output_columns=tuple(out_cols),
                rows=tuple(rows),
            ),
            [],
        )
