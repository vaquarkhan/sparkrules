from __future__ import annotations

from openpyxl import Workbook

from sparkrules.model.decision_table import ColumnType, DecisionTable


def _ct_label(c: ColumnType) -> str:
    return c.name


class DecisionTableExporter:
    @staticmethod
    def export(dt: DecisionTable, path: str) -> None:
        wb = Workbook()
        ws = wb.active
        if ws is None:
            raise RuntimeError("workbook has no active sheet")
        ws["A1"] = "RuleTable"
        ws["B1"] = "name"
        ws["C1"] = dt.name
        ws["A2"] = "HIT_POLICY"
        ws["B2"] = dt.hit_policy.name
        c = 1
        for col in dt.input_columns:
            ws.cell(3, c, "CONDITION")
            ws.cell(4, c, col.name)
            ws.cell(5, c, col.field_ref)
            ws.cell(6, c, _ct_label(col.col_type))
            ws.cell(7, c, col.operator or "==")
            c += 1
        for col in dt.output_columns:
            ws.cell(3, c, "ACTION")
            ws.cell(4, c, col.name)
            ws.cell(5, c, col.field_ref)
            ws.cell(6, c, _ct_label(col.col_type))
            ws.cell(7, c, "")
            c += 1
        n_in = len(dt.input_columns)
        n_out = len(dt.output_columns)
        for ri, r in enumerate(dt.rows, start=8):
            for j, v in enumerate(r.cells, start=1):
                ws.cell(ri, j, v)
            if r.priority != 0 or any(x.priority != 0 for x in dt.rows):
                ws.cell(ri, n_in + n_out + 1, r.priority)
        wb.save(path)
